"""Convert the newest Excel file in data/ into dist/data.js (used by the website)."""
import glob, json, os, shutil, sys, datetime as dt
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
files = [f for f in glob.glob(os.path.join(ROOT, "data", "*.xls*")) if not os.path.basename(f).startswith("~$")]
if not files:
    sys.exit("No Excel file found in data/")
src = max(files, key=os.path.getmtime)  # newest upload wins
print("Using:", os.path.basename(src))

wb = load_workbook(src, data_only=True, read_only=True)
DATE_COLS = {"PM due date", "CA due date", "PQ due date", "IQ date", "OQ date",
             "PQ Sign date", "PM sign", "CA sign", "PQ sign"}

def clean(h, v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if h in DATE_COLS and isinstance(v, (int, float)) and 20000 < v < 80000:  # Excel serial
        return (dt.date(1899, 12, 30) + dt.timedelta(days=int(v))).isoformat()
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return v if not isinstance(v, str) else v.strip()

rows = wb["EQ"].iter_rows(values_only=True)
headers = [str(h).strip() if h else None for h in next(rows)]
last = headers.index(None) if None in headers else len(headers)
headers = headers[:last]
items = []
for r in rows:
    it = {h: clean(h, r[i] if i < len(r) else None) for i, h in enumerate(headers)}
    if it.get("Equipment"):
        items.append({k: v for k, v in it.items() if v is not None})

people = {}
try:
    for r in wb["Personel"].iter_rows(min_row=2, values_only=True):
        if r[3] and r[4] and str(r[3]).strip().lower() != "abrv.":
            people[str(r[3]).strip()] = [str(r[4]).strip(), r[5]]
except KeyError:
    pass

out = os.path.join(ROOT, "dist")
os.makedirs(out, exist_ok=True)
for f in glob.glob(os.path.join(ROOT, "web", "*")):
    shutil.copy(f, out)
payload = {"updated": dt.datetime.fromtimestamp(os.path.getmtime(src)).strftime("%Y-%m-%d"),
           "source": os.path.basename(src), "headers": headers, "people": people, "items": items}
with open(os.path.join(out, "data.js"), "w", encoding="utf-8") as fh:
    fh.write("window.EQ_DATA=" + json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + ";")
print(f"{len(items)} equipment rows, {len(people)} people -> dist/data.js")
